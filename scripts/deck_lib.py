"""Editorial-deck library: design tokens + PPTX shape/text primitives.

Coordinates passed to every helper are *figure fractions* on a 16:9 canvas:
x runs 0 (left) -> 1 (right), y runs 0 (bottom) -> 1 (top).
"""
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE

EMU_W = 12192000          # 13.333 in
EMU_H = 6858000           # 7.5 in

SERIF = "Georgia"
SANS = "Arial"

# ---- design tokens (Elegant Editorial) ----
PAPER = "FAF8F3"          # warm paper background
INK = "2C2A24"             # near-black warm ink
SUB = "6E675C"             # muted warm grey
GOLD = "A98B4E"            # antique gold (rules, brackets)
GREEN = "33574B"           # deep moss (accent)
HAIR = "D8D0C0"            # warm hairline
DOT = "CFC7B6"             # dotted row rules
VDOT = "E3DCCB"            # dotted column rules
PALETTE = ["42607A", "5F7E6C", "A67C52", "8E5B52"]  # slate, sage, clay, rose
# extended muted hues for extra series
PALETTE_EXT = PALETTE + ["7A6E8A", "B08D3E", "4E7A78", "9C6B4E"]

DASH = {
    "dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
    "dash": MSO_LINE_DASH_STYLE.DASH,
    "dashdot": MSO_LINE_DASH_STYLE.DASH_DOT,
    "longdash": MSO_LINE_DASH_STYLE.LONG_DASH,
}

def C(h):
    return RGBColor.from_string(h)

def X(f): return int(f * EMU_W)
def Y(f): return int((1 - f) * EMU_H)
def W(f): return int(f * EMU_W)
def H(f): return int(f * EMU_H)

def kill_shadow(shape):
    shape.shadow.inherit = False

def new_deck(bg=PAPER):
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C(bg)
    return prs, slide

def hrect(slide, x, y, w, h, color, line=None, lw=1.0):
    """Filled rectangle; (x,y) = top-left corner."""
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, X(x), Y(y), W(w), H(h))
    s.fill.solid(); s.fill.fore_color.rgb = C(color)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = C(line); s.line.width = Pt(lw)
    kill_shadow(s)
    return s

def rrect(slide, x, y, w, h, color, adj=0.08, line=None, lw=1.0):
    """Rounded rectangle; (x,y) = centre. adj = corner radius / short side."""
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x - w/2),
                               Y(y + h/2), W(w), H(h))
    s.adjustments[0] = adj
    s.fill.solid(); s.fill.fore_color.rgb = C(color)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = C(line); s.line.width = Pt(lw)
    kill_shadow(s)
    return s

def txt(slide, x, y, text, size, color, bold=False, italic=False, family=SERIF,
        ha=PP_ALIGN.LEFT, va=MSO_ANCHOR.MIDDLE, bw=0.40, bh=0.045,
        wrap=False, spacing=None):
    """Text box with (x,y) as the true anchor point.

    bw, bh = box size in fractions; the box is positioned according to ha so
    that x is the left / centre / right anchor. Multi-line text uses '\\n'.
    """
    if ha == PP_ALIGN.CENTER:
        left = X(x - bw/2)
    elif ha == PP_ALIGN.RIGHT:
        left = X(x - bw)
    else:
        left = X(x)
    tb = slide.shapes.add_textbox(left, Y(y + bh/2), W(bw), H(bh))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = wrap
    tf.vertical_anchor = va
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ha
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.name = family; r.font.color.rgb = C(color)
        if spacing is not None:
            r._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
    return tb

def hline(slide, x0, x1, y, color, lw=1.0, dash=None):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(x0), Y(y), X(x1), Y(y))
    cn.line.color.rgb = C(color); cn.line.width = Pt(lw)
    if dash:
        cn.line.dash_style = DASH.get(dash, dash)
    kill_shadow(cn)
    return cn

def vline(slide, x, y0, y1, color, lw=1.0, dash=None):
    """Vertical line x fixed; y0 = bottom, y1 = top."""
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(x), Y(y0), X(x), Y(y1))
    cn.line.color.rgb = C(color); cn.line.width = Pt(lw)
    if dash:
        cn.line.dash_style = DASH.get(dash, dash)
    kill_shadow(cn)
    return cn

def end_tick(slide, x, y, bar_h, color, ext=0.012, lw_frac=0.0016):
    """Small vertical cap at the end of a bar (editorial bar terminator)."""
    return hrect(slide, x, y + bar_h/2 + ext, lw_frac, bar_h + 2*ext, color)

def bracket(slide, x0, x1, y, color, tick=0.012, lw=0.9):
    """Simple editorial bracket: horizontal rule with end ticks (pointing up)."""
    hline(slide, x0, x1, y, color, lw=lw)
    for xx in (x0, x1):
        hrect(slide, xx, y + tick, 0.0012, tick, color)
